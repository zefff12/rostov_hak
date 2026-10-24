"""Разбор исходных документов на структурные единицы (пункты, строки матрицы).

Единица поиска — пункт документа: так ответ может сослаться на конкретный
«п. 10.3», а не на абстрактный кусок текста.
"""
import re
from pathlib import Path

import pymupdf
from docx import Document

FOOTER_RE = re.compile(r"УЧЕБНЫЙ ПРИМЕР ДЛЯ ИТ-СОРЕВНОВАНИЙ.*", re.I)
POINT_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})(?:\.(\d{1,2}))?\.(?:\s|$)")
CHAPTER_RE = re.compile(r"^Глава\s+(\d+)\.\s*(.*)")
APPENDIX_RE = re.compile(r"^Приложение\s*№\s*(\d+)")
MAX_CHUNK = 1400


def _norm_ws(text: str) -> str:
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return text.strip()


def _join_lines(lines: list[str]) -> str:
    """Склеивает строки PDF в абзацы, сохраняя переносы перед маркерами списков."""
    out = ""
    for ln in lines:
        ln = ln.strip()
        if not ln:
            continue
        if not out:
            out = ln
        elif ln.startswith(("•", "—")) or re.match(r"^\d{1,2}\)", ln):
            out += "\n" + ln
        elif out.endswith("-") and not out.endswith(" -"):
            out += ln  # перенос по слогам / составное слово
        else:
            out += " " + ln
    return _norm_ws(re.sub(r" *\n *", "\n", out))


def split_long(text: str, limit: int = MAX_CHUNK) -> list[str]:
    """Делит длинный пункт на части по границам строк/предложений."""
    if len(text) <= limit:
        return [text]
    parts, cur = [], ""
    for piece in re.split(r"(?<=[;.:])\s+|\n", text):
        if cur and len(cur) + len(piece) + 1 > limit:
            parts.append(cur.strip())
            cur = ""
        cur += (" " if cur else "") + piece
    if cur.strip():
        parts.append(cur.strip())
    return parts


def parse_numbered(lines_with_pages: list[tuple[str, int]]) -> list[dict]:
    """Режет текст на пункты вида N.N. / N.N.N. с главами и приложениями.

    Номер пункта принимается только если он относится к текущей главе и идет
    по возрастанию — так ссылки «по пункту 2.13.» в начале строки не ломают разбор.
    """
    units: list[dict] = []
    chapter, chapter_title, last = None, "", (0, 0, 0)
    cur = None
    pending_title = False

    def close():
        nonlocal cur
        if cur:
            cur["text"] = _join_lines(cur.pop("lines"))
            if cur["text"]:
                units.append(cur)
        cur = None

    for line, page in lines_with_pages:
        s = line.strip()
        if not s:
            continue
        m_ch = CHAPTER_RE.match(s)
        m_app = APPENDIX_RE.match(s)
        m_pt = POINT_RE.match(s)
        if m_ch:
            close()
            chapter, chapter_title, last = int(m_ch.group(1)), m_ch.group(2).strip(), (int(m_ch.group(1)), 0, 0)
            pending_title = True
            continue
        if m_app:
            close()
            chapter, chapter_title, pending_title = None, f"Приложение № {m_app.group(1)}", False
            cur = {"ref": f"Приложение № {m_app.group(1)}", "section": chapter_title, "page": page, "lines": []}
            continue
        if m_pt and chapter is not None:
            num = (int(m_pt.group(1)), int(m_pt.group(2)), int(m_pt.group(3) or 0))
            if num[0] == chapter and num > last:
                close()
                pending_title = False
                last = num
                ref = ".".join(str(n) for n in num if n) if num[2] else f"{num[0]}.{num[1]}"
                cur = {"ref": f"п. {ref}", "section": f"Глава {chapter}. {chapter_title}", "page": page,
                       "lines": [s[m_pt.end():]]}
                continue
        if pending_title:
            chapter_title += " " + s
            continue
        if cur is not None:
            cur["lines"].append(s)
    close()
    return units


def pdf_lines(path: Path) -> list[tuple[str, int]]:
    doc = pymupdf.open(path)
    out = []
    for i, page in enumerate(doc, start=1):
        for ln in page.get_text().splitlines():
            s = ln.strip()
            if not s or FOOTER_RE.match(s) or re.fullmatch(r"\d{1,3}", s):
                continue
            out.append((s, i))
    return out


def docx_paragraphs(path: Path) -> list[str]:
    """Текст DOCX в порядке следования, включая ячейки таблиц."""
    doc = Document(path)
    out = []
    body = doc.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            t = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t"))
            if t.strip():
                out.append(_norm_ws(t))
        elif tag == "tbl":
            for row in child.iter():
                if not row.tag.endswith("}tr"):
                    continue
                cells = []
                for tc in row.iter():
                    if tc.tag.endswith("}tc"):
                        cells.append(_norm_ws(" ".join(
                            "".join(n.text or "" for n in p.iter() if n.tag.endswith("}t"))
                            for p in tc.iter() if p.tag.endswith("}p"))))
                if any(cells):
                    out.append(" | ".join(cells))
    return out


def parse_generic(path: Path) -> list[dict]:
    """Разбор произвольного документа: пункты, если они есть, иначе — абзацы."""
    if path.suffix.lower() == ".pdf":
        lines = pdf_lines(path)
    elif path.suffix.lower() == ".docx":
        lines = [(p, 1) for p in docx_paragraphs(path)]
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")
        lines = [(ln, 1) for ln in text.splitlines()]

    units = parse_numbered(lines)
    if len(units) < 3:
        units = _parse_simple_points(lines)
    if len(units) < 3:
        units = _parse_paragraphs(lines)
    return units


def _parse_simple_points(lines):
    """Пункты вида «1.1.» без глав (методические рекомендации и т.п.)."""
    units, cur, section = [], None, ""
    for s, page in lines:
        m = re.match(r"^(\d{1,2}\.\d{1,2}(?:\.\d{1,2})?)\.\s*(.*)", s)
        m_sec = re.match(r"^(\d{1,2})\.\s+([А-ЯЁ].{2,80})$", s)
        if m:
            if cur:
                units.append(cur)
            cur = {"ref": f"п. {m.group(1)}", "section": section, "page": page, "text": m.group(2)}
        elif m_sec and not (cur and cur["text"].rstrip().endswith(":")):
            if cur:
                units.append(cur)
                cur = None
            section = f"{m_sec.group(1)}. {m_sec.group(2)}"
        elif cur:
            cur["text"] += "\n" + s
    if cur:
        units.append(cur)
    return units


def _parse_paragraphs(lines):
    units, buf, page0 = [], "", 1
    for s, page in lines:
        if not buf:
            page0 = page
        buf += ("\n" if buf else "") + s
        if len(buf) > 800:
            units.append({"ref": f"фрагмент {len(units) + 1}", "section": "", "page": page0, "text": buf})
            buf = ""
    if buf:
        units.append({"ref": f"фрагмент {len(units) + 1}", "section": "", "page": page0, "text": buf})
    return units


# ---------------------------------------------------------------- Матрица

MATRIX_COLUMNS = ["Юридическое подразделение", "Департамент ДОУ", "Кадровое подразделение", "Прочие департаменты"]


def parse_matrix(path: Path) -> dict:
    """Матрица обязательного визирования → строки с отметками по столбцам."""
    doc = Document(path)
    preamble = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    rows, group = [], None
    for table in doc.tables:
        for r in table.rows:
            texts = [_norm_ws(c.text) for c in r.cells]
            first = texts[0]
            if first.upper() in ("ВНУТРЕННИЕ ДОКУМЕНТЫ", "ВНЕШНИЕ ДОКУМЕНТЫ"):
                group = "internal" if first.upper().startswith("ВНУТР") else "external"
                continue
            m = re.match(r"^(\d+)\.\s*(.*)", first, re.S)
            if not m or group is None:
                continue
            marks = texts[1:5]
            rows.append({
                "group": group,
                "num": int(m.group(1)),
                "title": m.group(2).strip(),
                "marks": dict(zip(MATRIX_COLUMNS, marks)),
            })
    return {"preamble": preamble, "rows": rows}


def matrix_units(matrix: dict) -> list[dict]:
    units = []
    for i, p in enumerate(matrix["preamble"], start=1):
        if re.match(r"^\d\.", p):
            units.append({"ref": f"общие положения, п. {p[0]}", "section": "Матрица обязательного визирования",
                          "page": 1, "text": p[2:].strip()})
    for r in matrix["rows"]:
        kind = "Внутренние документы" if r["group"] == "internal" else "Внешние документы"
        marks = "; ".join(f"{k}: {v or '—'}" for k, v in r["marks"].items())
        units.append({
            "ref": f"{kind.split()[0].lower()}, строка {r['num']}",
            "section": f"Матрица обязательного визирования — {kind}",
            "page": 1,
            "text": f"{r['title']}\nОбязательное визирование: {marks}",
        })
    return units
